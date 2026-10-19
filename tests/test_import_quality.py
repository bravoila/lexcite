"""Import/parsing reliability regressions (P1-8 ~ P1-14)."""
import hashlib
import io
import json
import zlib

import pytest
from app import db, ingest
from app.ingest import analyze_articles, extract_file, parse_articles


# ---------- P1-8 章节/分编行独立识别并从 body 剔除 ----------

def test_multiple_structural_headings_stripped_from_body():
    # An article can close a section AND a chapter at once: both heading lines
    # must leave the body, not just the first one.
    rows = parse_articles("第一条 正文一。\n第二节 测试节\n第二条 正文二。\n第二节 节二\n第二章 章二\n第三条 正文三。")
    assert [r["number"] for r in rows] == ["第一条", "第二条", "第三条"]
    assert rows[0]["body"] == "正文一。"
    assert rows[1]["body"] == "正文二。"


def test_part_heading_stripped():
    rows = parse_articles("第一条 正文一。\n第一部分 总则\n第二条 正文二。")
    assert rows[0]["body"] == "正文一。"


def test_heading_path_recorded_on_import(client):
    law = {"name": "测试章节目录", "version": "v1", "jurisdiction": "全国", "kind": "法律",
           "source_kind": "paste", "valid_from": "2021-01-01",
           "text": "第一章 总则\n第一条 正文一。"}
    lid = client.post("/api/laws", json=law).json()["id"]
    from app import indexing
    rows = [r for r in indexing.articles() if r["law_id"] == lid]
    assert rows[0]["heading_path"] == "第一章 总则"


# ---------- P1-9 引用启发式不吞真实条标题 ----------

def test_lead_in_line_with_consecutive_numbers_prefers_split():
    rows = parse_articles("第一条 本条的适用，参照《中华人民共和国合同法》\n第二条 测试条文二。\n第三条 测试条文三。")
    assert [r["number"] for r in rows] == ["第一条", "第二条", "第三条"]
    assert rows[1]["body"] == "测试条文二。"


def test_acceptance_lead_in_splits():
    rows = parse_articles("第一条 办理下列事项，依照下列规定\n第二条 测试条文二。")
    assert [r["number"] for r in rows] == ["第一条", "第二条"]


def test_ambiguous_reference_kept_with_warning():
    result = analyze_articles("第一条 正文参照《中华人民共和国合同法》\n第二条 应适用该法第二条的情形。")
    assert result["count"] == 1
    assert "第二条" in result["articles"][0]["body"]
    assert any("交叉引用" in w for w in result["warnings"])


# ---------- P1-10 SHA256 可审计(原文快照 + 重算) ----------

def _import_civil(client):
    law = json.loads((db.ROOT / "examples/civil_excerpt.json").read_text())
    r = client.post("/api/laws", json=law)
    assert r.status_code == 200, r.text
    return r.json()["id"], law


def test_snapshot_recompute_matches_stored_hash(client):
    lid, law = _import_civil(client)
    r = client.get(f"/api/laws/{lid}/hash-check")
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] and not body["snapshot_missing"]
    assert body["recomputed"] == body["sha256"] == hashlib.sha256(law["text"].encode()).hexdigest()
    assert "UTF-8" in body["note"]


def test_snapshot_is_compressed_original(client):
    lid, law = _import_civil(client)
    with db.connect() as c:
        blob = c.execute("SELECT content FROM law_snapshots WHERE law_id=?", (lid,)).fetchone()[0]
    assert zlib.decompress(blob).decode() == law["text"]


def test_hash_check_without_snapshot_reports_missing(client):
    lid, _ = _import_civil(client)
    with db.connect() as c:
        c.execute("DELETE FROM law_snapshots WHERE law_id=?", (lid,))
    r = client.get(f"/api/laws/{lid}/hash-check").json()
    assert r["snapshot_missing"] and not r["ok"]


# ---------- P1-11 PDF 扫描件拒绝 + DOCX 不丢内容 ----------

def _minimal_pdf(page_texts):
    """Hand-built single-font PDF (one page per entry, '' = blank), so the test
    does not depend on reportlab. ASCII only; pypdf extracts the Tj literal."""
    def stream(text):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        return b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(page_texts)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_texts)} >>".encode())
    for i, text in enumerate(page_texts):
        page = 3 + i * 2
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page + 1} 0 R /Resources << /Font << /F1 {3 + len(page_texts) * 2} 0 R >> >> >>".encode())
        objects.append(stream(text))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def test_scanned_pdf_rejected_with_clear_error():
    pdf = _minimal_pdf(["some readable provision text", "", ""])
    with pytest.raises(ValueError, match="扫描件"):
        extract_file("扫描.pdf", pdf)


def test_text_pdf_extracts_all_pages():
    pdf = _minimal_pdf(["some readable provision text", "another page of text"])
    text = extract_file("正文.pdf", pdf)
    assert "readable provision text" in text and "another page of text" in text


def test_docx_includes_tables_and_headers():
    from docx import Document
    doc = Document()
    doc.add_paragraph("第一条 正文一。")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "第二条 表格内条文。"
    doc.sections[0].header.paragraphs[0].text = "第三条 页眉条文。"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_file("正文.docx", buf.getvalue())
    assert "第一条 正文一。" in text and "第二条 表格内条文。" in text and "第三条 页眉条文。" in text
    assert text.index("第一条") < text.index("第二条")


def test_extract_endpoint_accepts_docx_off_event_loop(client):
    from docx import Document
    doc = Document()
    doc.add_paragraph("第一条 正文一。")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/api/extract", files={"file": ("正文.docx", buf.getvalue(),
                                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200 and "第一条 正文一。" in r.json()["text"]


# ---------- P1-13 重复导入查重返回已有记录 ----------

def test_duplicate_import_returns_existing_record(client, corpus):
    law = json.loads((db.ROOT / "examples/civil_excerpt.json").read_text())
    r = client.post("/api/laws", json=law)
    assert r.status_code == 200
    d = r.json()
    assert d["duplicate"] and d["id"] == corpus["law"]["id"] and "已存在法规" in d["message"]
    with db.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM laws").fetchone()[0] == 1


def test_same_text_different_metadata_imports_separately(client, corpus):
    # 相同正文但不同名称/版本是合法的独立版本(版本重叠复核场景),不得被查重误拦。
    law = json.loads((db.ROOT / "examples/civil_excerpt.json").read_text())
    law["name"] = "民法典别称测试"
    r = client.post("/api/laws", json=law)
    assert r.status_code == 200
    d = r.json()
    assert not d.get("duplicate") and d["id"] != corpus["law"]["id"]


# ---------- P1-14 Markdown 标记剥离 ----------

def test_markdown_markers_stripped_before_split():
    md = "# 第一条 正文一。\n* 第二条 正文二。\n- 第三条 正文三。"
    rows = parse_articles(extract_file("法规.md", md.encode()))
    assert [r["number"] for r in rows] == ["第一条", "第二条", "第三条"]


def test_txt_content_untouched():
    txt = "# 第一条 正文一。"
    assert extract_file("法规.txt", txt.encode()) == txt


def test_markdown_extraction_roundtrip_matches_plaintext_split():
    plain = "第一条 正文一。\n第二条 正文二。"
    md = "## 第一条 正文一。\n- 第二条 正文二。"
    assert parse_articles(extract_file("法规.md", md.encode())) == parse_articles(plain)


@pytest.mark.parametrize("marker", ["\ufeff", "\u200b", "\u200d", "\u200c", "\u2060"])
@pytest.mark.parametrize("separator", ["\u2028", "\u2029", "\x85", "\v", "\f", "\r\n"])
def test_invisible_markers_and_unicode_line_breaks_never_merge_articles(marker, separator):
    source = marker + "第一条 正文一。" + separator + marker + "第２条 正文二。" + separator + "第十 二条 正文十二。"
    result = analyze_articles(source)
    assert [ingest.article_key(row["number"]) for row in result["articles"]] == [(1, 0), (2, 0), (12, 0)]
    assert result["warnings"]


def test_normalized_parsing_keeps_exact_original_snapshot(client):
    source = "\ufeff第一条 原文一。\u2028\u200b第二条 原文二。"
    response = client.post("/api/laws", json={"name": "原文保真测试", "version": "v1", "valid_from": "2021-01-01", "text": source})
    assert response.status_code == 200 and response.json()["count"] == 2
    lid = response.json()["id"]
    with db.connect() as c:
        snapshot = c.execute("SELECT content FROM law_snapshots WHERE law_id=?", (lid,)).fetchone()[0]
    assert zlib.decompress(snapshot).decode() == source
    assert client.get(f"/api/laws/{lid}/hash-check").json()["sha256"] == hashlib.sha256(source.encode()).hexdigest()


def test_consecutive_number_does_not_turn_reference_into_fake_article():
    result = analyze_articles("第一条 依照《中华人民共和国合同法》\n第二条之规定，处理。\n第三条 正文三。")
    assert [row["number"] for row in result["articles"]] == ["第一条", "第三条"]
    assert "第二条之规定" in result["articles"][0]["body"]
    assert result["cross_reference_count"] == 1 and result["warnings"]


@pytest.mark.parametrize("heading", ["第一章 总则。", "第 一 章 总则;", "第一章 总则;附则", "第一章总则"])
def test_structural_heading_variants_share_parser_and_index_metadata(client, heading):
    text = "第一条 正文一。\n" + heading + "\n第二条 正文二。"
    result = analyze_articles(text)
    assert result["articles"][0]["body"] == "正文一。"
    assert result["heading_paths"]["第二条"] == heading
    lid = client.post("/api/laws", json={"name": "章节格式测试", "version": "v1", "valid_from": "2021-01-01", "text": text}).json()["id"]
    from app import indexing
    rows = [row for row in indexing.articles() if row["law_id"] == lid]
    assert rows[1]["heading_path"] == heading


def test_wrapped_chapter_reference_is_not_removed_or_indexed_as_heading():
    result = analyze_articles("第一章 总则\n第一条 应当依照\n第二章 规定的程序适用，依照前条规定。\n第二条 正文二。")
    assert "第二章 规定的程序适用" in result["articles"][0]["body"]
    assert result["heading_paths"]["第二条"] == "第一章 总则"


def test_rescued_heading_keeps_chapter_path_and_surfaces_import_warning(client):
    text = "第一章 总则\n第一条 依照《测试法规》\n第二条 正文二。\n第三条 正文三。"
    response = client.post("/api/laws", json={"name": "歧义提示测试", "version": "v1", "valid_from": "2021-01-01", "text": text})
    assert response.status_code == 200 and response.json()["warnings"]
    lid = response.json()["id"]
    from app import indexing
    assert all(row["heading_path"] == "第一章 总则" for row in indexing.articles() if row["law_id"] == lid)
    with db.connect() as c:
        audit = json.loads(c.execute("SELECT details FROM audit WHERE action='law_import'").fetchone()[0])
    assert audit["warnings"] == response.json()["warnings"]


def test_import_returns_reference_warnings_without_preview(client):
    text = "第一条 依据本法\n第九条的规定，处理相关事项。\n第二条 正文二。"
    response = client.post("/api/laws", json={"name": "引用提示测试", "version": "v1", "valid_from": "2021-01-01", "text": text})
    assert response.status_code == 200
    assert response.json()["cross_reference_count"] == 1 and response.json()["warnings"]


def test_duplicate_identity_is_checked_before_parsing(client, corpus):
    law = json.loads((db.ROOT / "examples/civil_excerpt.json").read_text())
    law["text"] = "这份文本没有任何合法条号，重复身份不应该重新解析。"
    response = client.post("/api/laws", json=law)
    assert response.status_code == 200 and response.json()["duplicate"]
    assert response.json()["id"] == corpus["law"]["id"]


def test_concurrent_duplicate_imports_return_same_record(client, monkeypatch):
    import threading
    from concurrent.futures import ThreadPoolExecutor
    from app.models import LawInput
    from app import indexing
    indexing.ensure_schema()
    barrier = threading.Barrier(2)
    original = ingest.analyze_articles
    def analyze(text):
        result = original(text)
        barrier.wait(timeout=3)
        return result
    monkeypatch.setattr(ingest, "analyze_articles", analyze)
    law = LawInput(name="并发重复测试", version="v1", valid_from="2021-01-01", text="第一条 并发导入测试正文。")
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(ingest.import_law, [law, law]))
    assert responses[0]["id"] == responses[1]["id"]
    assert sum(bool(response.get("duplicate")) for response in responses) == 1


def test_docx_recursive_structures_headers_and_merged_cells_are_not_lost_or_duplicated():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    doc = Document()
    doc.add_paragraph("正文段落。")
    table = doc.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "合并单元格正文。"
    merged.add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表格正文。"
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "内容控件正文。"
    run.append(text); paragraph.append(run); content.append(paragraph); control.append(content)
    doc.element.body.append(control)
    box = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "文本框正文。"
    run.append(text); paragraph.append(run); box.append(paragraph)
    doc.element.body.append(box)
    section = doc.sections[0]
    for index, container in enumerate((section.header, section.footer, section.first_page_header,
                                       section.first_page_footer, section.even_page_header, section.even_page_footer)):
        container.add_table(rows=1, cols=1, width=Inches(1)).cell(0, 0).text = f"页眉页脚表格{index}。"
    stream = io.BytesIO()
    doc.save(stream)
    result = extract_file("正文.docx", stream.getvalue())
    for expected in ("正文段落。", "合并单元格正文。", "嵌套表格正文。", "内容控件正文。", "文本框正文。",
                     *(f"页眉页脚表格{i}。" for i in range(6))):
        assert result.count(expected) == 1


def test_file_extraction_does_not_block_other_requests(client, monkeypatch):
    import threading
    from concurrent.futures import ThreadPoolExecutor
    entered, release = threading.Event(), threading.Event()
    def slow_extract(*args):
        entered.set()
        assert release.wait(3)
        return "第一条 测试正文。"
    monkeypatch.setattr(ingest, "extract_file", slow_extract)
    with ThreadPoolExecutor(max_workers=2) as pool:
        extraction = pool.submit(client.post, "/api/extract", files={"file": ("正文.docx", b"test")})
        try:
            assert entered.wait(1)
            assert pool.submit(client.get, "/healthz").result(timeout=1).status_code == 200
        finally:
            release.set()
        assert extraction.result(timeout=2).status_code == 200


@pytest.mark.parametrize("name,content", [("broken.pdf", b"%PDF-broken"), ("broken.docx", b"not-a-zip")])
def test_corrupt_documents_have_explicit_parse_errors(name, content):
    with pytest.raises(ValueError, match="文件损坏"):
        extract_file(name, content)


def test_pdf_has_aggregate_extracted_text_limit(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_EXTRACTED_CHARS", 35)
    with pytest.raises(ValueError, match="超过200万字符"):
        extract_file("正文.pdf", _minimal_pdf(["long legal provision text", "another long legal provision text"]))


@pytest.mark.parametrize("markdown", [
    "**第一条** 正文一。\n> 1. **第二条** 正文二。",
    "####### 第一条 正文一。\n> > - **第二条** 正文二。",
    "__第一条__ 正文一。\n1) 第二条 正文二。",
])
def test_nested_markdown_does_not_silently_drop_articles(markdown):
    assert parse_articles(extract_file("正文.md", markdown.encode())) == parse_articles("第一条 正文一。\n第二条 正文二。")


def test_txt_keeps_bom_while_markdown_removes_it():
    source = "\ufeff第一条 原文内容。"
    assert extract_file("正文.txt", source.encode()) == source
    assert extract_file("正文.md", source.encode()) == source.lstrip("\ufeff")


def test_mixed_and_fullwidth_article_digits_and_invalid_zero_suffix():
    assert ingest.article_key("第1百23条") == (123, 0)
    assert ingest.article_key("第１２３条之２") == (123, 2)
    for label in ("第1条之零", "第1条之0", "第零条"):
        with pytest.raises(ValueError, match="大于零"):
            ingest.article_key(label)


def test_corrupt_snapshot_reports_explicit_error(client):
    lid, _ = _import_civil(client)
    with db.connect() as c:
        c.execute("UPDATE law_snapshots SET content=? WHERE law_id=?", (b"corrupt", lid))
    response = client.get(f"/api/laws/{lid}/hash-check")
    assert response.status_code == 409 and "原文存档损坏" in response.json()["detail"]
